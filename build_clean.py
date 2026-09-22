#!/usr/bin/env python3
"""Clean, readable single-sheet wedding budget with a Paid/Not-paid status column.
All amounts in full Rupiah (IDR)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook(); ws = wb.active; ws.title = "Budget Pernikahan"
ws.sheet_view.showGridLines = False
F = "Calibri"
thin = Side(style="thin", color="D9D2C2")
med  = Side(style="medium", color="9B8E72")
GREEN="2E7D46"; RED="B23A3A"; AMBER="B07A1E"; INK="3B2F1B"; MUT="8A7F68"
GREY = PatternFill("solid", fgColor="EFEADF")     # section band
SUB  = PatternFill("solid", fgColor="F6F3EC")
TOT  = PatternFill("solid", fgColor="E7DFC9")
HEAD = PatternFill("solid", fgColor="3B2F1B")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
LTOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
CTR  = Alignment(horizontal="center", vertical="center", wrap_text=True)
RGT  = Alignment(horizontal="right", vertical="center")
RP   = '"Rp"#,##0'
row_h_only = Border(bottom=thin)   # light horizontal rule only -> not a rigid grid

for c,w in {"A":5,"B":46,"C":12,"D":18,"E":14,"F":40}.items():
    ws.column_dimensions[c].width = w

# title
ws["A1"]="BUDGET PERNIKAHAN  —  [ Mempelai Wanita ]  &  [ Mempelai Pria ]"
ws["A1"].font=Font(name=F,size=15,bold=True,color=INK); ws.merge_cells("A1:F1"); ws.row_dimensions[1].height=24
ws["A2"]="Adat Jawa · The Bimasena (Rooang), Dharmawangsa · semua angka dalam Rupiah"
ws["A2"].font=Font(name=F,size=10,italic=True,color=MUT); ws.merge_cells("A2:F2")

HR=4
for i,h in enumerate(["No","Rincian","Qty / Pax","Biaya","Status","Catatan"]):
    c=ws.cell(HR,i+1,h); c.font=Font(name=F,size=10,bold=True,color="FFFFFF"); c.fill=HEAD
    c.alignment=CTR if i in (0,2,3,4) else Alignment(horizontal="left",vertical="center")
    c.border=Border(bottom=med)
ws.row_dimensions[HR].height=20

# status -> font colour
def stcolor(s):
    s=s.lower()
    if s.startswith("lunas"): return GREEN
    if "dp" in s: return AMBER
    return RED

sections=[
 ("A · LAMARAN / TUNANGAN  —  November 2026",[
   ("Tenda + kursi (sewa)","",1500000,"Belum",""),
   ("Dekorasi & bunga","",1500000,"Belum",""),
   ("Tambahan makanan & minuman","",1000000,"Belum","Masak sendiri di rumah"),
   ("Seserahan / hantaran","",500000,"Belum","Nampan + wrapping"),
   ("Dokumentasi","½ hari",300000,"Belum",""),
   ("Makeup mempelai wanita (lamaran)","",2000000,"DP dibayar","DP sudah dibayar"),
   ("Buffer / lain-lain","",200000,"Belum",""),
 ]),
 ("B · SIRAMAN  —  H-1",[
   ("Upacara siraman","",3000000,"Belum","Bunga setaman, air 7 sumber, kendi, dawet, tumpeng, gendhing"),
   ("Katering keluarga (siraman)","",1500000,"Belum","Makan keluarga di rumah"),
 ]),
 ("C · HARI-H  —  AKAD & PANGGIH",[
   ("The Bimasena (Rooang) — minimum spending","100+ pax",150000000,"Belum","Min. spend; sudah termasuk F&B 100+ pax"),
   ("Perlengkapan panggih (temu manten)","",2000000,"Belum","Kembar mayang ×2, suruh, telur, kacar-kucur, sindur"),
   ("Penghulu / KUA (akad)","",600000,"Belum","Di luar jam kantor"),
   ("Koordinator hari-H (WO on-the-day)","",3000000,"Belum","Jaga rundown, cue vendor"),
   ("MC / pranatacara (akad + resepsi)","",2500000,"Belum","Bilingual"),
 ]),
 ("D · RESEPSI",[
   ("Dekorasi & bunga resepsi","",5000000,"Belum","Melati putih: pelaminan, entrance, aisle, meja"),
   ("Hiburan / live music","",3500000,"Belum","Trio akustik"),
   ("Kamar hotel — bridal suite (1 malam)","",2500000,"Belum","Malam pernikahan"),
   ("Wedding cake & dessert","",1000000,"Belum",""),
   ("Souvenir / favours","100 pcs",1500000,"Belum","± 100 × 15rb"),
   ("Foto & video","2 hari",6000000,"Belum","Siraman + hari-H"),
 ]),
 ("E · BUSANA, RIAS & CINCIN",[
   ("Cincin tunangan + cincin nikah","",23000000,"Lunas","Total kedua cincin — sudah dibeli"),
   ("Busana & rias mempelai wanita","",10000000,"Belum","Paes ageng: siraman, akad, resepsi"),
   ("Busana mempelai pria — sewa beskap (akad)","",4000000,"Belum","Hanya beskap akad; tak perlu jas lamaran/resepsi"),
   ("Busana adat orang tua","",2000000,"Belum","Kebaya / beskap"),
 ]),
 ("F · LAIN-LAIN",[
   ("Undangan","",300000,"Belum","Digital (+ cetak bila perlu)"),
   ("Mahar","",1000000,"Belum","Simbolis — sesuai niat"),
   ("Transportasi & lain-lain","",2000000,"Belum","Logistik 2 hari, mobil, tip"),
 ]),
]

r=HR+1; n=0; sub_rows=[]; venue_row=rings_row=None; first_item=r
for title,items in sections:
    ws.cell(r,1,title).font=Font(name=F,size=11,bold=True,color=INK)
    for cc in range(1,7): ws.cell(r,cc).fill=GREY
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
    ws.cell(r,1).alignment=Alignment(horizontal="left",vertical="center",indent=1)
    ws.row_dimensions[r].height=22; r+=1
    first=r
    for label,qty,amt,status,note in items:
        n+=1
        ws.cell(r,1,n).font=Font(name=F,size=10,color=MUT); ws.cell(r,1).alignment=CTR
        ws.cell(r,2,label).font=Font(name=F,size=10,color=INK); ws.cell(r,2).alignment=LTOP
        ws.cell(r,3,qty).font=Font(name=F,size=10,color=MUT); ws.cell(r,3).alignment=CTR
        d=ws.cell(r,4,amt); d.font=Font(name=F,size=10,color=INK); d.number_format=RP; d.alignment=RGT
        e=ws.cell(r,5,status); e.font=Font(name=F,size=9,bold=True,color=stcolor(status)); e.alignment=CTR
        ws.cell(r,6,note).font=Font(name=F,size=9,color=MUT); ws.cell(r,6).alignment=LTOP
        for cc in range(1,7): ws.cell(r,cc).border=row_h_only
        if "bimasena" in label.lower(): venue_row=r
        if label.startswith("Cincin"): rings_row=r
        ws.row_dimensions[r].height=28; r+=1
    last=r-1
    ws.cell(r,2,f"Subtotal {title[0]}").font=Font(name=F,size=10,bold=True,italic=True,color=MUT); ws.cell(r,2).alignment=RGT
    s=ws.cell(r,4,f"=SUM(D{first}:D{last})"); s.font=Font(name=F,size=10,bold=True,color=INK); s.number_format=RP; s.alignment=RGT
    for cc in range(1,7): ws.cell(r,cc).fill=SUB; ws.cell(r,cc).border=Border(top=thin)
    sub_rows.append(r); ws.row_dimensions[r].height=18; r+=2
last_item=last

# contingency (10% of non-fixed: exclude venue & paid rings)
n+=1
ws.cell(r,1,n).font=Font(name=F,size=10,color=MUT); ws.cell(r,1).alignment=CTR
ws.cell(r,2,"Biaya tak terduga (10%)").font=Font(name=F,size=10,color=INK); ws.cell(r,2).alignment=LTOP
subsum="+".join(f"D{x}" for x in sub_rows)
d=ws.cell(r,4,f"=0.1*(({subsum})-D{venue_row}-D{rings_row})"); d.font=Font(name=F,size=10,color=INK); d.number_format=RP; d.alignment=RGT
ws.cell(r,5,"Belum").font=Font(name=F,size=9,bold=True,color=RED); ws.cell(r,5).alignment=CTR
ws.cell(r,6,"10% pos non-fix (di luar venue & cincin)").font=Font(name=F,size=9,color=MUT); ws.cell(r,6).alignment=LTOP
for cc in range(1,7): ws.cell(r,cc).border=row_h_only
cont=r; r+=2

def total(r,label,formula,big=False,fill=TOT):
    ws.cell(r,2,label).font=Font(name=F,size=12 if big else 11,bold=True,color=INK); ws.cell(r,2).alignment=RGT
    t=ws.cell(r,4,formula); t.font=Font(name=F,size=13 if big else 11,bold=True,color=INK); t.number_format=RP; t.alignment=RGT
    for cc in range(1,7): ws.cell(r,cc).fill=fill; ws.cell(r,cc).border=Border(top=med,bottom=med)
    ws.row_dimensions[r].height=24 if big else 20

total(r,"TOTAL ANGGARAN",f"=({subsum})+D{cont}",big=True); gr=r; r+=1
total(r,"Sudah dibayar (Lunas)",f'=SUMIF(E{first_item}:E{last_item},"Lunas",D{first_item}:D{last_item})',fill=SUB); paid=r
ws.cell(r,4).font=Font(name=F,size=11,bold=True,color=GREEN); r+=1
total(r,"Sisa yang belum dibayar",f"=D{gr}-D{paid}",fill=SUB)
ws.cell(r,4).font=Font(name=F,size=11,bold=True,color=RED); r+=2

# funding note
ws.cell(r,1,"CATATAN PENDANAAN").font=Font(name=F,size=11,bold=True,color=INK)
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6)
for cc in range(1,7): ws.cell(r,cc).fill=GREY
ws.cell(r,1).alignment=Alignment(horizontal="left",vertical="center",indent=1); r+=1
for label,val,note in [
   ("Saldo Blugether saat ini (Sep 2026)",50000000,""),
   ("Menabung per bulan",5000000,"Wanita 2 jt + Pria 3 jt"),
   ("Sudah dibayar di luar tabungan",None,"Cincin Rp 23 jt (lunas) + DP makeup lamaran"),
]:
    ws.cell(r,2,label).font=Font(name=F,size=10,color=INK); ws.cell(r,2).alignment=LTOP
    if val is not None:
        c=ws.cell(r,4,val); c.font=Font(name=F,size=10,color=INK); c.number_format=RP; c.alignment=RGT
    ws.cell(r,6,note).font=Font(name=F,size=9,color=MUT); ws.cell(r,6).alignment=LTOP
    ws.row_dimensions[r].height=16; r+=1

ws.freeze_panes="A5"
out="/Users/gomobile/Documents/Project/Wedding-Budget/Budget_Pernikahan_Bimasena.xlsx"
wb.save(out); print("saved", out)
